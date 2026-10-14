import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from unicodedata import combining, normalize

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..main import ClassRoom, Guardian, GuardianStudent, Student, audit, error, require

router = APIRouter(tags=['Imports'])
MAX_IMPORT_ROWS = 2000
MAX_XLSX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 1000
IMPORT_FIELDS = ('nis', 'full_name', 'class_name', 'major', 'gender', 'guardian_name', 'guardian_phone')
REQUIRED_FIELDS = ('nis', 'full_name', 'class_name')
FIELD_LABELS = {
    'nis': 'NIS/NISN',
    'full_name': 'Nama siswa',
    'class_name': 'Kelas/Rombel',
    'major': 'Jurusan',
    'gender': 'Jenis kelamin',
    'guardian_name': 'Nama wali',
    'guardian_phone': 'Nomor wali',
}
HEADER_ALIASES = {
    'nis': ('nis', 'nisn', 'nis nisn', 'nomor induk siswa', 'nomor induk', 'student id'),
    'full_name': ('nama', 'nama siswa', 'nama lengkap', 'nama lengkap siswa', 'name', 'full name'),
    'class_name': ('kelas', 'rombel', 'kelas rombel', 'class', 'class name'),
    'major': ('jurusan', 'jurusan program keahlian', 'program keahlian', 'kompetensi keahlian', 'program', 'major'),
    'gender': ('gender', 'jenis kelamin', 'jk'),
    'guardian_name': ('nama wali', 'nama orang tua', 'nama wali murid'),
    'guardian_phone': ('nomor wali', 'nomor orang tua', 'telepon wali', 'no hp wali', 'nomor hp wali'),
}
PHONE_PATTERN = re.compile(r'(\+62|62|0)\d{8,13}')


def _error(status: int, message: str, code: str, data=None):
    detail = {'success': False, 'message': message, 'errors': {}, 'code': code}
    if data is not None:
        detail['data'] = data
    raise HTTPException(status, detail)


async def _read_upload(file: UploadFile) -> tuple[str, bytes]:
    suffix = Path(file.filename or '').suffix.lower()
    if suffix not in ('.csv', '.xlsx'):
        _error(415, 'File harus berformat .xlsx atau .csv.', 'IMPORT_FILE_TYPE')
    raw = await file.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if len(raw) > settings.max_upload_mb * 1024 * 1024:
        _error(413, 'Ukuran file melebihi batas.', 'IMPORT_FILE_TOO_LARGE')
    if not raw:
        _error(422, 'File kosong.', 'IMPORT_FILE_EMPTY')
    return suffix, raw


def _cell_text(value, number_format: str = 'General') -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int):
        text = str(value)
        if re.fullmatch(r'0+', number_format) and len(number_format) > len(text):
            return text.zfill(len(number_format))
        return text
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
        if re.fullmatch(r'0+', number_format) and len(number_format) > len(text):
            return text.zfill(len(number_format))
        return text
    return str(value).strip()


def _parse_csv(raw: bytes) -> tuple[list[str], list[tuple[int, list[str]]]]:
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        _error(422, 'CSV harus menggunakan encoding UTF-8.', 'IMPORT_ENCODING')
    reader = csv.reader(io.StringIO(text, newline=''))
    try:
        headers = next(reader, None)
    except csv.Error:
        _error(422, 'Struktur CSV tidak valid.', 'IMPORT_FILE_INVALID')
    if headers is None or not any(str(value).strip() for value in headers):
        _error(422, 'File tidak memiliki header.', 'IMPORT_FILE_EMPTY')
    if len(headers) > 200:
        _error(422, 'File memiliki terlalu banyak kolom (maksimal 200).', 'IMPORT_TOO_MANY_COLUMNS')
    rows = []
    try:
        for row_number, values in enumerate(reader, start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                _error(422, f'File melebihi batas {MAX_IMPORT_ROWS} baris data.', 'IMPORT_TOO_MANY_ROWS')
            if any(str(value).strip() for value in values):
                rows.append((row_number, [str(value).strip() for value in values]))
    except csv.Error:
        _error(422, 'Struktur CSV tidak valid atau memiliki field yang terlalu besar.', 'IMPORT_FILE_INVALID')
    return [str(value).strip() for value in headers], rows


def _parse_xlsx(raw: bytes) -> tuple[list[str], list[tuple[int, list[str]]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES or sum(item.file_size for item in entries) > MAX_XLSX_EXPANDED_BYTES:
                _error(422, 'Isi file Excel melebihi batas pemrosesan.', 'IMPORT_FILE_TOO_LARGE')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError):
        _error(422, 'File Excel tidak valid atau tidak dapat dibaca.', 'IMPORT_FILE_INVALID')
    try:
        worksheet = workbook.active
        if worksheet.max_column > 200:
            _error(422, 'File memiliki terlalu banyak kolom (maksimal 200).', 'IMPORT_TOO_MANY_COLUMNS')
        if worksheet.max_row > MAX_IMPORT_ROWS + 1:
            _error(422, f'File melebihi batas {MAX_IMPORT_ROWS} baris data.', 'IMPORT_TOO_MANY_ROWS')
        first = next(worksheet.iter_rows(min_row=1, max_row=1), ())
        headers = [_cell_text(cell.value, cell.number_format) for cell in first]
        if not headers or not any(headers):
            _error(422, 'File tidak memiliki header.', 'IMPORT_FILE_EMPTY')
        rows = []
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                _error(422, f'File melebihi batas {MAX_IMPORT_ROWS} baris data.', 'IMPORT_TOO_MANY_ROWS')
            values = [_cell_text(cell.value, cell.number_format) for cell in cells]
            if any(values):
                rows.append((row_number, values))
        return headers, rows
    finally:
        workbook.close()


def _normalized_header(value: str) -> str:
    decomposed = normalize('NFKD', value.casefold())
    return re.sub(r'[^a-z0-9]+', ' ', ''.join(ch for ch in decomposed if not combining(ch))).strip()


def _suggest_mapping(headers: list[str]) -> dict[str, int]:
    aliases = {
        field: {_normalized_header(alias) for alias in names}
        for field, names in HEADER_ALIASES.items()
    }
    mapping = {}
    claimed = set()
    for index, header in enumerate(headers):
        normalized = _normalized_header(header)
        matches = [field for field, candidates in aliases.items() if normalized in candidates]
        if len(matches) == 1 and matches[0] not in mapping and index not in claimed:
            mapping[matches[0]] = index
            claimed.add(index)
    return mapping


def _resolve_mapping(headers: list[str], mapping_json: str) -> dict[str, int]:
    if not mapping_json:
        mapping = _suggest_mapping(headers)
    else:
        try:
            submitted = json.loads(mapping_json)
        except json.JSONDecodeError:
            _error(422, 'Pemetaan kolom tidak valid.', 'IMPORT_MAPPING_INVALID')
        if not isinstance(submitted, dict):
            _error(422, 'Pemetaan kolom tidak valid.', 'IMPORT_MAPPING_INVALID')
        mapping = {}
        for field, index in submitted.items():
            if field not in IMPORT_FIELDS or index in (None, ''):
                continue
            if isinstance(index, bool) or not isinstance(index, int) or index < 0 or index >= len(headers):
                _error(422, 'Pemetaan menunjuk kolom yang tidak tersedia.', 'IMPORT_MAPPING_INVALID')
            mapping[field] = index
    if len(set(mapping.values())) != len(mapping):
        _error(422, 'Satu kolom tidak dapat dipetakan ke lebih dari satu field.', 'IMPORT_MAPPING_INVALID')
    return mapping


def _row_value(values: list[str], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    return values[index].strip() if index is not None and index < len(values) else ''


def _preview_result(headers, raw_rows, mapping, db: Session) -> dict:
    classes = db.scalars(select(ClassRoom).where(ClassRoom.is_active.is_(True))).all()
    classes_by_name: dict[str, list[ClassRoom]] = {}
    for classroom in classes:
        classes_by_name.setdefault(classroom.name.casefold().strip(), []).append(classroom)
    nis_candidates = {
        _row_value(values, mapping, 'nis')
        for _, values in raw_rows
        if _row_value(values, mapping, 'nis')
    }
    existing_nis = set()
    candidate_list = list(nis_candidates)
    for offset in range(0, len(candidate_list), 500):
        existing_nis.update(db.scalars(
            select(Student.nis).where(Student.nis.in_(candidate_list[offset:offset + 500]))
        ).all())
    seen_nis = set()
    rows = []
    mapped_required = [field for field in REQUIRED_FIELDS if field not in mapping]

    for row_number, values in raw_rows:
        data = {field: _row_value(values, mapping, field) for field in IMPORT_FIELDS}
        errors = []
        duplicate = False
        nis = data['nis']
        if nis and (nis in existing_nis or nis in seen_nis):
            duplicate = True
            errors.append({'field': 'nis', 'message': 'NIS sudah terdaftar atau duplikat dalam file.'})
        elif nis:
            seen_nis.add(nis)

        for field in mapped_required:
            errors.append({'field': field, 'message': f'Kolom {FIELD_LABELS[field]} belum dipetakan.'})
        if not mapped_required:
            if not nis:
                errors.append({'field': 'nis', 'message': 'NIS wajib diisi.'})
            elif len(nis) > 20:
                errors.append({'field': 'nis', 'message': 'NIS maksimal 20 karakter.'})
            full_name = data['full_name']
            if len(full_name) < 2:
                errors.append({'field': 'full_name', 'message': 'Nama siswa minimal 2 karakter.'})
            elif len(full_name) > 120:
                errors.append({'field': 'full_name', 'message': 'Nama siswa maksimal 120 karakter.'})
            class_name = data['class_name'].casefold().strip()
            major = data['major'].casefold().strip()
            candidates = classes_by_name.get(class_name, [])
            if not class_name:
                errors.append({'field': 'class_name', 'message': 'Kelas/Rombel wajib diisi.'})
            elif major:
                candidates = [item for item in candidates if item.major.casefold().strip() == major]
                if not candidates:
                    errors.append({'field': 'class_name', 'message': 'Kelas dan jurusan aktif tidak ditemukan.'})
            elif not candidates:
                errors.append({'field': 'class_name', 'message': 'Kelas aktif tidak ditemukan.'})
            elif len(candidates) > 1:
                errors.append({'field': 'major', 'message': 'Nama kelas digunakan beberapa jurusan; petakan kolom jurusan.'})
            elif candidates:
                data['class_id'] = str(candidates[0].id)
            if major and len(candidates) == 1:
                data['class_id'] = str(candidates[0].id)
            if data['gender']:
                gender_value = data['gender'].casefold()
                if gender_value in ('l', 'laki-laki', 'laki laki', 'male', 'm'):
                    data['gender'] = 'L'
                elif gender_value in ('p', 'perempuan', 'female', 'f'):
                    data['gender'] = 'P'
                else:
                    errors.append({'field': 'gender', 'message': 'Jenis kelamin harus L/P, laki-laki, atau perempuan.'})
            if mapping.get('guardian_name') is not None and mapping.get('guardian_phone') is not None:
                guardian_name, guardian_phone = data['guardian_name'], data['guardian_phone']
                if bool(guardian_name) != bool(guardian_phone):
                    errors.append({'field': 'guardian_name', 'message': 'Nama dan nomor wali harus diisi bersama.'})
                elif guardian_name:
                    if len(guardian_name) < 2:
                        errors.append({'field': 'guardian_name', 'message': 'Nama wali minimal 2 karakter.'})
                    if not PHONE_PATTERN.fullmatch(guardian_phone):
                        errors.append({'field': 'guardian_phone', 'message': 'Nomor telepon Indonesia tidak valid.'})

        if duplicate:
            status = 'duplicate'
        elif errors:
            status = 'invalid'
        else:
            status = 'valid'
        rows.append({
            'row_number': row_number,
            'status': status,
            'valid': status == 'valid',
            'errors': errors,
            'data': data,
        })

    counts = {
        'total_rows': len(rows),
        'valid_rows': sum(row['status'] == 'valid' for row in rows),
        'duplicate_rows': sum(row['status'] == 'duplicate' for row in rows),
        'invalid_rows': sum(row['status'] == 'invalid' for row in rows),
    }
    return {
        **counts,
        'headers': [{'index': index, 'label': header or f'Kolom {index + 1}'} for index, header in enumerate(headers)],
        'mapping': mapping,
        'mapping_required': bool(mapped_required),
        'rows': rows,
    }


async def _parse_upload(file: UploadFile, mapping_json: str, db: Session):
    suffix, raw = await _read_upload(file)
    if suffix == '.csv':
        headers, rows = _parse_csv(raw)
    else:
        headers, rows = _parse_xlsx(raw)
    if not rows:
        _error(422, 'File tidak berisi data siswa.', 'IMPORT_FILE_EMPTY')
    if len(headers) > 200:
        _error(422, 'File memiliki terlalu banyak kolom (maksimal 200).', 'IMPORT_TOO_MANY_COLUMNS')
    mapping = _resolve_mapping(headers, mapping_json)
    return _preview_result(headers, rows, mapping, db)


@router.get('/api/students/import-template.csv')
def import_template(u=Depends(require('ADMIN_IT'))):
    return StreamingResponse(
        iter(['nis,nama,kelas,jurusan,nama_wali,nomor_wali\n']),
        media_type='text/csv',
        headers={'Content-Disposition': 'attachment; filename="tandara-template-siswa.csv"'},
    )


@router.post('/api/students/import/preview')
async def import_preview(
    file: UploadFile = File(...),
    mapping: str = Form(default=''),
    db: Session = Depends(get_db),
    u=Depends(require('ADMIN_IT')),
):
    return {'success': True, 'data': await _parse_upload(file, mapping, db)}


@router.post('/api/students/import')
async def import_students(
    file: UploadFile = File(...),
    mapping: str = Form(default=''),
    duplicate_policy: str = Form(default='skip'),
    db: Session = Depends(get_db),
    u=Depends(require('ADMIN_IT')),
):
    if duplicate_policy not in ('skip', 'abort'):
        _error(422, 'Kebijakan duplikasi tidak valid.', 'IMPORT_POLICY_INVALID')
    result = await _parse_upload(file, mapping, db)
    if duplicate_policy == 'abort' and result['duplicate_rows']:
        _error(409, 'Import dibatalkan karena ditemukan NIS duplikat.', 'IMPORT_DUPLICATES_FOUND', {
            'imported_count': 0,
            'skipped_count': 0,
            'failed_count': result['invalid_rows'],
            'duplicate_rows': result['duplicate_rows'],
            'rows': result['rows'],
        })
    valid_rows = [row for row in result['rows'] if row['status'] == 'valid']
    created = 0
    try:
        class_ids = {int(row['data']['class_id']) for row in valid_rows}
        active_class_ids = set()
        if class_ids:
            active_class_ids = set(db.scalars(
                select(ClassRoom.id).where(ClassRoom.id.in_(class_ids), ClassRoom.is_active.is_(True))
            ).all())
        if active_class_ids != class_ids:
            _error(409, 'Data kelas berubah setelah preview. Tinjau ulang file sebelum import.', 'IMPORT_DATA_CHANGED')

        pending_links = []
        for row in valid_rows:
            data = row['data']
            guardian = None
            if data['guardian_name'] and data['guardian_phone']:
                guardian = Guardian(
                    full_name=data['guardian_name'],
                    phone_number=data['guardian_phone'],
                )
                db.add(guardian)
            student = Student(
                nis=data['nis'],
                full_name=data['full_name'],
                class_id=int(data['class_id']),
                gender=data['gender'] or None,
                guardian=guardian,
            )
            db.add(student)
            if guardian is not None:
                pending_links.append((guardian, student))
            created += 1
        db.flush()
        db.add_all(
            GuardianStudent(guardian_id=guardian.id, student_id=student.id, relationship='Wali')
            for guardian, student in pending_links
        )
        if created:
            audit(db, u, 'IMPORT', 'Student', None, f'Mengimpor {created} siswa')
        db.commit()
    except IntegrityError:
        db.rollback()
        _error(409, 'Data berubah selama import (kemungkinan NIS baru saja terdaftar). Tidak ada baris yang disimpan.', 'IMPORT_CONFLICT')
    except Exception:
        db.rollback()
        raise

    return {'success': True, 'data': {
        'imported_count': created,
        'importedCount': created,
        'skipped_count': result['duplicate_rows'],
        'failed_count': result['invalid_rows'],
        'duplicate_rows': result['duplicate_rows'],
        'rows': result['rows'],
    }}
